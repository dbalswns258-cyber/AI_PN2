"""Deterministic, review-first DXF takeoff. No web or database dependencies."""
import csv
import hashlib
import json
import math
from collections import Counter, defaultdict
from decimal import Decimal, ROUND_DOWN, ROUND_HALF_EVEN, ROUND_HALF_UP
from pathlib import Path

import ezdxf
from openpyxl import Workbook

from . import ENGINE_VERSION

UNIT_FACTORS = {'mm': Decimal('0.001'), 'cm': Decimal('0.01'), 'm': Decimal('1')}
INSUNITS = {4: 'mm', 5: 'cm', 6: 'm'}
ROUNDINGS = {'ROUND_HALF_UP': ROUND_HALF_UP, 'ROUND_HALF_EVEN': ROUND_HALF_EVEN, 'ROUND_DOWN': ROUND_DOWN}
MAX_ENTITIES = 20000
MAX_VERTICES = 50000


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, allow_nan=False).encode()).hexdigest()


def file_digest(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(65536), b''):
            h.update(chunk)
    return h.hexdigest()


def _point(v):
    p = [float(v[0]), float(v[1]), float(v[2])]
    if not all(math.isfinite(x) for x in p):
        raise ValueError('좌표에 유효하지 않은 숫자가 있습니다.')
    return p


def _segments(points, bulges, closed):
    segments = []
    for i in range(len(points) if closed else len(points) - 1):
        a, b = points[i], points[(i + 1) % len(points)]
        bulge = float(bulges[i])
        if not math.isfinite(bulge):
            raise ValueError('유효하지 않은 폴리선 곡률입니다.')
        chord = math.hypot(b[0] - a[0], b[1] - a[1])
        if chord == 0 and bulge != 0:
            raise ValueError('시작과 끝이 같은 곡선 구간은 지원하지 않습니다.')
        length = chord if bulge == 0 else chord * (1 + bulge * bulge) / (4 * abs(bulge)) * 4 * math.atan(abs(bulge))
        # Bulge is signed sagitta / half chord. Midpoint distinguishes circular arcs.
        midpoint = [(a[0] + b[0]) / 2 + (b[1] - a[1]) * bulge / 2,
                    (a[1] + b[1]) / 2 - (b[0] - a[0]) * bulge / 2,
                    (a[2] + b[2]) / 2]
        if not math.isfinite(length):
            raise ValueError('길이가 계산 범위를 초과했습니다.')
        segments.append({'start': a, 'end': b, 'bulge': bulge, 'length_model': str(length), 'midpoint': midpoint})
    return segments


def _geometry(e):
    kind = e.dxftype()
    if kind == 'LINE':
        points = [_point(e.dxf.start), _point(e.dxf.end)]
        bulges, closed = [0, 0], False
    elif kind == 'LWPOLYLINE':
        if len(e) > MAX_VERTICES:
            raise ValueError('객체별 정점 수 제한을 초과했습니다.')
        if tuple(e.dxf.extrusion) != (0, 0, 1):
            raise ValueError('기울어진 객체 좌표계(OCS)는 지원하지 않습니다.')
        raw = list(e.get_points('xyb'))
        points = [[float(x), float(y), float(e.dxf.elevation)] for x, y, _ in raw]
        bulges, closed = [b for _, _, b in raw], e.closed
    elif kind == 'POLYLINE':
        if not e.is_2d_polyline:
            raise ValueError('3D 폴리선·메시·폴리페이스는 미지원입니다.')
        if e.dxf.flags & 6:
            raise ValueError('곡선 맞춤/스플라인 맞춤 POLYLINE은 미지원입니다.')
        if tuple(e.dxf.extrusion) != (0, 0, 1):
            raise ValueError('기울어진 객체 좌표계(OCS)는 지원하지 않습니다.')
        if len(e.vertices) > MAX_VERTICES:
            raise ValueError('객체별 정점 수 제한을 초과했습니다.')
        points = [_point(v.dxf.location) for v in e.vertices]
        for p in points:
            p[2] += float(e.dxf.elevation.z)
        bulges, closed = [v.dxf.get('bulge', 0) for v in e.vertices], e.is_closed
    else:
        raise ValueError('길이 산출 미지원 객체입니다.')
    if len(points) < 2:
        raise ValueError('정점이 2개 미만입니다.')
    if not all(math.isfinite(v) for p in points for v in p):
        raise ValueError('유효하지 않은 좌표입니다.')
    if max(p[2] for p in points) - min(p[2] for p in points) > 1e-9:
        raise ValueError('높이가 변하는 객체는 첫 2D 산출 범위에서 제외합니다.')
    segments = _segments(points, bulges, closed)
    length = sum((Decimal(s['length_model']) for s in segments), Decimal(0))
    if length <= 0:
        raise ValueError('길이가 0인 객체입니다.')
    return points, segments, str(length)


def _geometry_key(segments):
    # Exact coincident segments, independent of entity direction and polyline start vertex.
    # No tolerance-based merging: close, partly overlapping and disconnected lines stay separate.
    keys = []
    for s in segments:
        endpoints = sorted([tuple(s['start']), tuple(s['end'])])
        keys.append((endpoints[0], endpoints[1], tuple(s['midpoint'])))
    return digest(sorted(keys))


def extract(path):
    doc = ezdxf.readfile(path)
    model = doc.modelspace()
    if len(model) > MAX_ENTITIES:
        raise ValueError(f'첫 버전은 모델 공간 객체 {MAX_ENTITIES}개까지 지원합니다.')
    entities, duplicates = [], defaultdict(list)
    for e in model:
        item = {'handle': e.dxf.handle, 'type': e.dxftype(), 'layer': e.dxf.layer,
                'color': e.dxf.color, 'linetype': e.dxf.linetype, 'supported': False,
                'length_model': None, 'points': [], 'segments': [], 'warnings': [], 'text': '', 'block': None}
        if e.dxftype() in ('LINE', 'LWPOLYLINE', 'POLYLINE'):
            try:
                item['points'], item['segments'], item['length_model'] = _geometry(e)
                item['supported'] = True
                duplicates[_geometry_key(item['segments'])].append(item['handle'])
            except ValueError as exc:
                item['warnings'].append(str(exc))
        elif e.dxftype() in ('TEXT', 'MTEXT'):
            item['text'] = e.dxf.text if e.dxftype() == 'TEXT' else e.plain_text()
            item['points'] = [_point(e.dxf.insert)]
            item['warnings'].append('문자는 판독 참고 정보이며 길이 산출 대상이 아닙니다.')
        elif e.dxftype() == 'INSERT':
            block = doc.blocks.get(e.dxf.name)
            item['block'] = {'name': e.dxf.name, 'insert': _point(e.dxf.insert),
                             'scale': [e.dxf.xscale, e.dxf.yscale, e.dxf.zscale], 'rotation': e.dxf.rotation,
                             'attributes': [{'tag': a.dxf.tag, 'text': a.dxf.text} for a in e.attribs],
                             'xref': bool(block and block.block.dxf.flags & 12),
                             'children': dict(Counter(x.dxftype() for x in block)) if block else {}}
            item['points'] = [_point(e.dxf.insert)]
            item['warnings'].append('블록 내부·배열·외부참조는 펼치거나 산출하지 않습니다. 별도 검토가 필요합니다.')
        else:
            item['warnings'].append(f'{e.dxftype()} 객체는 미지원이며 산출에서 제외됩니다.')
        entities.append(item)
    groups = [sorted(handles) for handles in duplicates.values() if len(handles) > 1]
    for item in entities:
        item['duplicate_handles'] = next((g for g in groups if item['handle'] in g), [])
        if item['duplicate_handles']:
            item['warnings'].append('완전히 겹치는 객체입니다. 하나를 선택하고 나머지를 명시적으로 제외하세요.')
    paper_count = sum(len(layout) for layout in doc.layouts if layout.name != 'Model')
    return {'engine_version': ENGINE_VERSION, 'sha256': file_digest(path),
            'insunits': doc.header.get('$INSUNITS', 0), 'suggested_unit': INSUNITS.get(doc.header.get('$INSUNITS', 0)),
            'entities': sorted(entities, key=lambda x: x['handle']), 'duplicate_groups': groups,
            'warnings': [f'배치 공간 객체 {paper_count}개는 길이 산출에서 제외됩니다.',
                         '부분 겹침·미세 중복·끊긴 선·평면 교차의 연결 관계는 자동 판정하지 않습니다.',
                         '객체 하나는 하나의 관경만 지정합니다. 관경 변경 구간은 CAD에서 분할 후 새 버전으로 등록하세요.',
                         '블록 내부와 외부참조를 펼치지 않습니다. TEXT/MTEXT는 참고 정보만 추출합니다.']}


def validate_decision(entity, decision):
    status = decision.get('status', 'pending')
    if status not in ('pending', 'confirmed', 'excluded'):
        raise ValueError('잘못된 확인 상태입니다.')
    if status != 'pending' and not decision.get('evidence', '').strip():
        raise ValueError('확정·제외에는 판단 근거가 필요합니다.')
    if status == 'confirmed':
        if not entity['supported']:
            raise ValueError('미지원 객체는 배관으로 확정할 수 없습니다.')
        if decision.get('system') != 'water':
            raise ValueError('첫 버전에서는 급수만 확정할 수 있습니다.')
        try:
            dn = Decimal(str(decision.get('diameter_mm', '')))
        except Exception as exc:
            raise ValueError('호칭지름을 입력하세요.') from exc
        if not dn.is_finite() or dn <= 0 or dn > 10000:
            raise ValueError('호칭지름은 0 초과 10000 이하의 숫자여야 합니다.')


def calculate(extraction, decisions, unit, unit_evidence, rules):
    if unit not in UNIT_FACTORS or not unit_evidence.strip():
        raise ValueError('모델 공간 길이 단위와 확인 근거가 필요합니다.')
    if rules.get('length_basis') != 'model_xy_centerline' or rules.get('round_at') != 'diameter_total':
        raise ValueError('지원하지 않는 길이 또는 반올림 위치 기준입니다.')
    for field in ('fitting_adjustment', 'waste'):
        if rules.get(field) != 'not_applied':
            raise ValueError('부속 조정·할증은 아직 지원하지 않습니다.')
    if rules.get('riser_and_fixture_allowances') != 'not_calculated' or rules.get('insulation') != 'not_calculated':
        raise ValueError('입상·기구 연결 보완과 보온은 아직 지원하지 않습니다.')
    if rules.get('quantity_origin') != 'drawing_direct' or not rules.get('version'):
        raise ValueError('직접 도면 물량 기준과 버전이 필요합니다.')
    places = rules.get('decimal_places')
    if type(places) is not int or not 0 <= places <= 6 or rules.get('rounding') not in ROUNDINGS:
        raise ValueError('반올림 설정이 잘못되었습니다.')
    handles = {e['handle'] for e in extraction['entities']}
    if set(decisions) - handles:
        raise ValueError('원본에 없는 객체 확인 정보가 포함되어 있습니다.')
    factor = UNIT_FACTORS[unit]
    totals, rows = defaultdict(Decimal), []
    for entity in extraction['entities']:
        d = decisions.get(entity['handle'], {'status': 'pending'})
        validate_decision(entity, d)
        status, reason = d.get('status', 'pending'), ''
        if status == 'confirmed' and any(decisions.get(other, {}).get('status') != 'excluded'
                                         for other in entity['duplicate_handles'] if other != entity['handle']):
            status, reason = 'pending', '중복 객체 중 하나를 선택하고 나머지를 제외해야 합니다.'
        length = None
        diameter = str(Decimal(str(d['diameter_mm'])).normalize()) if d.get('diameter_mm') else None
        if diameter:
            diameter = format(Decimal(diameter), 'f')
        if status == 'confirmed':
            length = Decimal(entity['length_model']) * factor
            totals[diameter] += length
        rows.append({'handle': entity['handle'], 'type': entity['type'], 'layer_original': entity['layer'],
                     'status': status, 'decision_status': d.get('status', 'pending'), 'system': d.get('system'),
                     'diameter_mm': diameter, 'material': d.get('material') or None, 'floor': d.get('floor') or None,
                     'zone': d.get('zone') or None, 'work_category': d.get('work_category') or None,
                     'quantity_origin': 'drawing_direct' if length is not None else 'unresolved_or_excluded',
                     'length_model': entity['length_model'], 'length_m': str(length) if length is not None else None,
                     'formula': f'{entity["length_model"]} × {factor}' if length is not None else '',
                     'evidence': d.get('evidence', ''), 'warnings': entity['warnings'] + ([reason] if reason else []),
                     'location': entity['points'][:1], 'rule_version': rules['version']})
    quantum = Decimal(1).scaleb(-places)
    summary = [{'diameter_mm': diameter, 'raw_length_m': str(total),
                'length_m': str(total.quantize(quantum, rounding=ROUNDINGS[rules['rounding']]))}
               for diameter, total in sorted(totals.items(), key=lambda pair: Decimal(pair[0]))]
    return {'engine_version': ENGINE_VERSION, 'source_sha256': extraction['sha256'], 'unit': unit,
            'unit_evidence': unit_evidence, 'rules': rules, 'rules_sha256': digest(rules),
            'decisions_sha256': digest(decisions), 'rows': rows, 'summary': summary,
            'counts': dict(Counter(row['status'] for row in rows)),
            'warnings': extraction['warnings'] + ['평면 직접 물량만 계산합니다. 입상·기구 연결·부속·할증·보온은 미산출입니다.',
                                                '미확정 또는 제외 항목이 있으면 집계는 도면 전체 물량을 뜻하지 않습니다.']}


def _cell(value):
    # Force all user-originated strings to text; never write spreadsheet formulas.
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    return value


def _sheet(workbook, name, records, fields):
    sheet = workbook.create_sheet(name)
    sheet.append(fields)
    for record in records:
        sheet.append([_cell(record.get(field)) for field in fields])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(55, max(16, len(str(column[0].value)) + 3))
    return sheet


def export(source, output_dir, extraction, result):
    if file_digest(source) != result['source_sha256'] or extraction['sha256'] != result['source_sha256']:
        raise ValueError('원본 파일 해시가 다릅니다. 기존 결과와 혼합할 수 없습니다.')
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=False)
    doc = ezdxf.readfile(source)
    doc.appids.new('AI_PN2') if 'AI_PN2' not in doc.appids else None
    standard_layers = {}
    for row in result['rows']:
        if row['status'] != 'confirmed':
            continue
        diameter = row['diameter_mm']
        if diameter not in standard_layers:
            base = 'STD_WATER_DN' + diameter.replace('.', '_')
            layer, number = base, 1
            while layer in doc.layers:
                layer = f'{base}_TAKEOFF_{number}'
                number += 1
            doc.layers.new(layer, dxfattribs={'color': 5, 'linetype': 'Continuous'})
            standard_layers[diameter] = layer
        layer = standard_layers[diameter]
        e = doc.entitydb[row['handle']]
        e.dxf.layer, e.dxf.color, e.dxf.linetype = layer, 256, 'BYLAYER'
        e.dxf.discard('true_color')
        e.set_xdata('AI_PN2', [(1000, f'source_sha256={result["source_sha256"]}'),
                                (1000, f'source_handle={row["handle"]}'),
                                (1000, f'rule={result["rules"]["version"]}')])
    doc.saveas(output / 'standardized.dxf')
    (output / 'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, '관경별집계', result['summary'], ['diameter_mm', 'raw_length_m', 'length_m'])
    fields = ['handle', 'type', 'layer_original', 'status', 'decision_status', 'system', 'diameter_mm', 'material',
              'floor', 'zone', 'work_category', 'quantity_origin', 'length_model', 'length_m', 'formula', 'evidence',
              'location', 'warnings', 'rule_version']
    _sheet(workbook, '객체별근거', result['rows'], fields)
    _sheet(workbook, '미확정및제외', [r for r in result['rows'] if r['status'] != 'confirmed'], fields)
    metadata = [{ 'key': k, 'value': result[k]} for k in ['engine_version', 'source_sha256', 'unit', 'unit_evidence',
                                                        'rules', 'rules_sha256', 'decisions_sha256', 'counts', 'warnings']]
    _sheet(workbook, '실행기준', metadata, ['key', 'value'])
    if result.get('provenance'):
        _sheet(workbook, '프로젝트원본연결', [{'key': k, 'value': v} for k, v in result['provenance'].items()], ['key', 'value'])
    if 'comparison' in result:
        _sheet(workbook, '기존산출비교', result['comparison'], ['handle', 'diameter_mm', 'baseline_m', 'calculated_m', 'difference_m', 'status'])
    workbook.save(output / 'takeoff.xlsx')


def compare_baseline(result, baseline_path):
    """Explicit segment handle + nominal diameter mapping; tolerance is not invented."""
    with open(baseline_path, encoding='utf-8-sig', newline='') as f:
        reader = csv.DictReader(f)
        if set(reader.fieldnames or []) != {'handle', 'diameter_mm', 'length_m'}:
            raise ValueError('비교 CSV 열은 handle,diameter_mm,length_m 이어야 합니다.')
        baseline = {}
        for row in reader:
            diameter, length = Decimal(row['diameter_mm']), Decimal(row['length_m'])
            if not diameter.is_finite() or diameter <= 0 or not length.is_finite() or length < 0:
                raise ValueError('비교 수치가 유효하지 않습니다.')
            key = (row['handle'], format(diameter.normalize(), 'f'))
            if key in baseline:
                raise ValueError('비교 자료에 중복 구간이 있습니다.')
            baseline[key] = length
    actual = {(r['handle'], r['diameter_mm']): Decimal(r['length_m']) for r in result['rows'] if r['length_m'] is not None}
    rows = []
    for key in sorted(set(baseline) | set(actual)):
        before, after = baseline.get(key), actual.get(key)
        rows.append({'handle': key[0], 'diameter_mm': key[1], 'baseline_m': str(before) if before is not None else None,
                     'calculated_m': str(after) if after is not None else None,
                     'difference_m': str(after - before) if before is not None and after is not None else None,
                     'status': 'matched' if before is not None and after is not None else 'missing_counterpart'})
    return rows
