import copy
import json
import math
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
import ezdxf
from openpyxl import load_workbook
from takeoff.engine import calculate, compare_baseline, export, extract, file_digest
from takeoff.sample import create_sample


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = create_sample(self.root / 'sample')
        self.extraction = extract(self.source)
        self.decisions = json.loads((self.source.parent / 'decisions.json').read_text())
        self.rules = json.loads((Path(__file__).resolve().parents[1] / 'rules/planar-v1.json').read_text())

    def result(self, **kwargs):
        args = dict(extraction=self.extraction, decisions=self.decisions, unit='mm', unit_evidence='합성 치수 검증', rules=self.rules)
        args.update(kwargs)
        return calculate(**args)

    def test_known_line_polyline_arc_totals(self):
        result = self.result()
        self.assertEqual({r['diameter_mm']: r['length_m'] for r in result['summary']}, {'20': '8.142', '25': '3.000'})
        self.assertEqual(result['counts'], {'confirmed': 3, 'excluded': 1, 'pending': 5})
        self.assertAlmostEqual(sum(float(r['length_m']) for r in result['rows'] if r['diameter_mm'] == '20'), 5 + math.pi)

    def test_deterministic_result(self):
        self.assertEqual(self.result(), self.result(decisions=dict(reversed(list(self.decisions.items())))))

    def test_unknown_units_and_missing_evidence_rejected(self):
        for override in ({'unit': ''}, {'unit': 'inch'}, {'unit_evidence': ' '}):
            with self.assertRaises(ValueError):
                self.result(**override)

    def test_explicit_unit_conversion(self):
        self.assertEqual(self.result(unit='m')['summary'][1]['length_m'], '3000.000')

    def test_print_scale_does_not_change_model_length(self):
        doc = ezdxf.readfile(self.source)
        doc.header['$DIMSCALE'] = 100
        other = self.root / 'scale.dxf'; doc.saveas(other)
        self.assertEqual(self.result(extraction=extract(other))['summary'], self.result()['summary'])

    def test_duplicate_requires_explicit_exclusion(self):
        decisions = {k: v for k, v in self.decisions.items() if v['status'] != 'excluded'}
        result = self.result(decisions=decisions)
        self.assertEqual(result['summary'][0]['length_m'], '3.142')
        self.assertTrue(any('중복 객체 중' in ' '.join(r['warnings']) for r in result['rows']))

    def test_double_confirmed_duplicates_are_both_blocked(self):
        decisions = copy.deepcopy(self.decisions)
        group = self.extraction['duplicate_groups'][0]
        for h in group:
            decisions[h] = {'status': 'confirmed', 'system': 'water', 'diameter_mm': '20', 'evidence': 'test'}
        result = self.result(decisions=decisions)
        self.assertTrue(all(r['status'] == 'pending' for r in result['rows'] if r['handle'] in group))

    def test_unsupported_block_circle_and_vertical_stay_unresolved(self):
        for entity in self.extraction['entities']:
            if entity['type'] in ('INSERT', 'CIRCLE') or '높이가 변하는' in ' '.join(entity['warnings']):
                self.assertFalse(entity['supported'])
                decisions = {entity['handle']: {'status': 'confirmed', 'system': 'water', 'diameter_mm': '20', 'evidence': 'test'}}
                with self.assertRaises(ValueError):
                    self.result(decisions=decisions)

    def test_unknown_handle_and_missing_evidence_rejected(self):
        with self.assertRaises(ValueError):
            self.result(decisions={'UNKNOWN': {'status': 'pending'}})
        decisions = copy.deepcopy(self.decisions)
        next(iter(decisions.values()))['evidence'] = ''
        with self.assertRaises(ValueError):
            self.result(decisions=decisions)

    def test_closed_polyline_includes_closing_segment(self):
        doc = ezdxf.new(); doc.modelspace().add_lwpolyline([(0, 0), (3, 0), (3, 4)], close=True)
        path = self.root / 'closed.dxf'; doc.saveas(path)
        self.assertEqual(Decimal(extract(path)['entities'][0]['length_model']), Decimal('12'))

    def test_legacy_polyline_bulge(self):
        doc = ezdxf.new(); p = doc.modelspace().add_polyline2d([(0, 0), (2000, 0)])
        p.vertices[0].dxf.bulge = 1
        path = self.root / 'legacy.dxf'; doc.saveas(path)
        self.assertAlmostEqual(float(extract(path)['entities'][0]['length_model']), math.pi * 1000)

    def test_tilted_ocs_is_not_projected_silently(self):
        doc = ezdxf.new(); p = doc.modelspace().add_lwpolyline([(0, 0), (10, 0)])
        p.dxf.extrusion = (0, 1, 0)
        path = self.root / 'tilt.dxf'; doc.saveas(path)
        self.assertFalse(extract(path)['entities'][0]['supported'])

    def test_partial_overlap_is_not_silently_merged(self):
        doc = ezdxf.new(); m = doc.modelspace()
        m.add_line((0, 0), (10, 0)); m.add_line((5, 0), (15, 0))
        path = self.root / 'overlap.dxf'; doc.saveas(path)
        self.assertEqual(extract(path)['duplicate_groups'], [])
        self.assertIn('부분 겹침', ' '.join(extract(path)['warnings']))

    def test_export_preserves_source_handles_and_other_entities(self):
        before = self.source.read_bytes(); result = self.result()
        export(self.source, self.root / 'out', self.extraction, result)
        self.assertEqual(self.source.read_bytes(), before)
        doc = ezdxf.readfile(self.root / 'out/standardized.dxf')
        self.assertEqual(len(doc.modelspace()), len(self.extraction['entities']))
        for row in result['rows']:
            entity = doc.entitydb[row['handle']]
            if row['status'] == 'confirmed':
                self.assertEqual(entity.dxf.layer, 'STD_WATER_DN' + row['diameter_mm'])
                self.assertIn(('source_handle=' + row['handle']), [x.value for x in entity.get_xdata('AI_PN2')])
            else:
                self.assertEqual(entity.dxf.layer, row['layer_original'])
        with self.assertRaises(FileExistsError):
            export(self.source, self.root / 'out', self.extraction, result)

    def test_source_hash_mismatch_rejected(self):
        result = self.result(); result['source_sha256'] = '0' * 64
        with self.assertRaises(ValueError):
            export(self.source, self.root / 'out', self.extraction, result)

    def test_excel_contains_exclusions_and_literal_formula_text(self):
        decisions = copy.deepcopy(self.decisions)
        next(iter(decisions.values()))['evidence'] = '=HYPERLINK("https://example.com")'
        result = self.result(decisions=decisions)
        export(self.source, self.root / 'out', self.extraction, result)
        wb = load_workbook(self.root / 'out/takeoff.xlsx')
        self.assertEqual(wb['미확정및제외'].max_row, 7)
        cells = [c for row in wb['객체별근거'] for c in row if isinstance(c.value, str) and c.value.startswith('=HYPERLINK')]
        self.assertEqual(len(cells), 1); self.assertEqual(cells[0].data_type, 's')
        self.assertIn('source_sha256', [r[0].value for r in wb['실행기준']])

    def test_rounding_only_after_diameter_total(self):
        doc = ezdxf.new(); m = doc.modelspace()
        a = m.add_line((0, 0), (0.4, 0)); b = m.add_line((1, 0), (1.4, 0))
        path = self.root / 'round.dxf'; doc.saveas(path)
        decisions = {e.dxf.handle: {'status': 'confirmed', 'system': 'water', 'diameter_mm': '20', 'evidence': 'test'} for e in (a, b)}
        result = self.result(extraction=extract(path), decisions=decisions)
        self.assertEqual(result['summary'][0]['length_m'], '0.001')

    def test_unsupported_business_adjustment_rejected(self):
        rules = {**self.rules, 'waste': '5%'}
        with self.assertRaises(ValueError):
            self.result(rules=rules)

    def test_baseline_compares_handle_and_diameter_without_assumed_tolerance(self):
        result = self.result(); row = next(r for r in result['rows'] if r['status'] == 'confirmed')
        path = self.root / 'baseline.csv'
        path.write_text(f'handle,diameter_mm,length_m\n{row["handle"]},{row["diameter_mm"]},4.9\nmissing,25,1.0\n')
        comparison = compare_baseline(result, path)
        matched = next(r for r in comparison if r['handle'] == row['handle'])
        self.assertEqual(Decimal(matched['difference_m']), Decimal('0.1'))
        self.assertTrue(any(r['status'] == 'missing_counterpart' for r in comparison))

    def test_layer_name_collision_and_truecolor_do_not_override_standard(self):
        doc = ezdxf.readfile(self.source)
        doc.layers.new('STD_WATER_DN20', dxfattribs={'color': 1})
        handle = next(iter(self.decisions))
        doc.entitydb[handle].dxf.true_color = 0xFF0000
        doc.saveas(self.source)
        extraction = extract(self.source); result = self.result(extraction=extraction)
        export(self.source, self.root / 'out', extraction, result)
        output = ezdxf.readfile(self.root / 'out/standardized.dxf')
        self.assertEqual(output.layers.get('STD_WATER_DN20').dxf.color, 1)
        self.assertFalse(output.entitydb[handle].dxf.hasattr('true_color'))
        self.assertEqual(output.entitydb[handle].dxf.layer, 'STD_WATER_DN20_TAKEOFF_1')
